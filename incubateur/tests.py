"""Tests des fonctions essentielles : python manage.py test"""
import datetime as dt
import io
import os
import tempfile
from unittest import mock

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse

from .models import Commentaire, CompteRendu, Document, Jalon, Projet, Suivi, User, cohorte_courante

DONNEES = {"title": "Test projet", "domain": "TIC, numérique, IA et données", "problem": "pb", "solution": "sol"}


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class ParcoursTests(TestCase):
    def setUp(self):
        def mk(e, r):
            return User.objects.create_user(username=e, email=e, password="Passw0rd!", first_name=e, role=r)
        self.porteur, self.comite = mk("p@x.sn", "porteur"), mk("c@x.sn", "comite")
        self.coordo, self.mentor, self.autre = mk("co@x.sn", "coordo"), mk("m@x.sn", "mentor"), mk("m2@x.sn", "mentor")

    def test_acces_protege_et_csrf(self):
        self.assertEqual(self.client.get(reverse("projets")).status_code, 302)
        r = Client(enforce_csrf_checks=True).post(reverse("login"), {"username": "a", "password": "b"})
        self.assertEqual(r.status_code, 403)

    def test_depot_completude_et_droits(self):
        self.client.force_login(self.porteur)
        r = self.client.post(reverse("projet_nouveau"), DONNEES, follow=True)
        p = Projet.objects.get(title="Test projet")
        self.assertEqual(p.completude()[0], 80)
        self.assertContains(r, "80 %")
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 403)

    def test_fichier_interdit_refuse(self):
        self.client.force_login(self.porteur)
        data = dict(DONNEES, file=SimpleUploadedFile("virus.exe", b"x"))
        self.client.post(reverse("projet_nouveau"), data)
        self.assertEqual(Projet.objects.count(), 0)

    def test_parcours_complet(self):
        p = Projet.objects.create(owner=self.porteur, title="T", domain="TIC, numérique, IA et données", problem="a", solution="b")
        self.client.force_login(self.comite)
        self.client.post(reverse("evaluer", args=[p.pk]), {"innovation": 4, "faisabilite": 4, "impact": 5, "equipe": 3,
                                                           "decision": "accepte", "motivation": "ok"})
        p.refresh_from_db()
        self.assertEqual(p.status, "accepte")
        self.assertEqual(p.evaluations.first().total, 16)
        self.client.force_login(self.coordo)
        self.client.post(reverse("affecter", args=[p.pk]), {"mentor_id": self.mentor.pk})
        self.client.post(reverse("jalon_ajout", args=[p.pk]), {"title": "J1", "due": "2020-01-01"})
        p.refresh_from_db()
        self.assertEqual(p.mentor, self.mentor)
        self.assertTrue(Jalon.objects.get(title="J1").en_retard)
        self.assertContains(self.client.get(reverse("dashboard") + "?cohort=2026-A"), "Taux de sélection")
        self.assertEqual(self.client.get(reverse("export_xlsx")).content[:2], b"PK")
        self.assertEqual(self.client.get(reverse("fiche_pdf", args=[p.pk])).content[:4], b"%PDF")
        self.client.force_login(self.mentor)
        self.assertEqual(self.client.get(reverse("projet", args=[p.pk])).status_code, 200)
        self.client.force_login(self.autre)
        self.assertEqual(self.client.get(reverse("projet", args=[p.pk])).status_code, 403)

    def test_ia_suggestion_puis_validation(self):
        p = Projet.objects.create(owner=self.porteur, title="T", domain="TIC, numérique, IA et données")
        self.client.force_login(self.porteur)
        with mock.patch.dict(os.environ):
            os.environ.pop("ANTHROPIC_API_KEY", None)
            r = self.client.post(reverse("ia"), {"texte": "Mon idée", "mode": "pitch"})
        self.assertContains(r, "mode démonstration")
        self.client.post(reverse("ia_accepter"), {"texte": "Texte validé", "project_id": p.pk, "champ": "problem"})
        p.refresh_from_db()
        self.assertEqual((p.problem, p.ia_used), ("Texte validé", True))


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class SuiviTests(TestCase):
    """Objectifs, documents et preuves, commentaires, suivis 3/6 mois, calendrier, compte rendu, tableau de bord."""

    def setUp(self):
        def mk(e, r):
            return User.objects.create_user(username=e, email=e, password="Passw0rd!", first_name=e, role=r)
        self.porteur, self.coordo = mk("p@x.sn", "porteur"), mk("co@x.sn", "coordo")
        self.mentor, self.autre = mk("m@x.sn", "mentor"), mk("m2@x.sn", "mentor")
        self.p = Projet.objects.create(owner=self.porteur, title="Séchoir", domain="TIC, numérique, IA et données",
                                       status="accepte", mentor=self.mentor)
        self.j = Jalon.objects.create(project=self.p, title="Prototype", due=dt.date.today() - dt.timedelta(days=5))

    def test_objectifs_et_suivis(self):
        self.client.force_login(self.mentor)
        fin = dt.date.today() - dt.timedelta(days=100)
        self.client.post(reverse("accompagnement", args=[self.p.pk]), {"objectifs": "Vendre 10 unités",
                                                                       "fin_incubation": fin.isoformat()})
        self.p.refresh_from_db()
        self.assertEqual(self.p.objectifs, "Vendre 10 unités")
        self.assertEqual([m for m, d in self.p.suivis_dus()], [3])
        self.client.force_login(self.coordo)
        self.assertContains(self.client.get(reverse("dashboard")), "suivi à 3 mois")
        self.client.force_login(self.mentor)
        data = {"periode": 3, "situation": "active", "emplois": 2, "observations": "RAS"}
        self.client.post(reverse("suivi_ajout", args=[self.p.pk]), data)
        self.client.post(reverse("suivi_ajout", args=[self.p.pk]), data)  # doublon refusé
        self.assertEqual(Suivi.objects.filter(project=self.p).count(), 1)
        self.assertEqual(self.p.suivis_dus(), [])
        self.client.force_login(self.porteur)
        r = self.client.post(reverse("suivi_ajout", args=[self.p.pk]), dict(data, periode=6))
        self.assertEqual(r.status_code, 403)

    def test_documents_preuves_et_droits(self):
        self.client.force_login(self.porteur)
        self.client.post(reverse("document_ajout", args=[self.p.pk]),
                         {"titre": "Budget", "fichier": SimpleUploadedFile("budget.pdf", b"%PDF-1.4")})
        self.client.post(reverse("document_ajout", args=[self.p.pk]),
                         {"titre": "Photo", "jalon_id": self.j.pk, "fichier": SimpleUploadedFile("proto.jpg", b"x")})
        self.client.post(reverse("document_ajout", args=[self.p.pk]),
                         {"titre": "Virus", "fichier": SimpleUploadedFile("x.exe", b"x")})
        self.assertEqual(Document.objects.filter(jalon=None).count(), 1)
        self.assertEqual(self.j.preuves.count(), 1)
        self.assertNotIn("Pièce jointe", self.p.completude()[1])
        d = Document.objects.get(titre="Budget")
        self.assertEqual(self.client.get(reverse("document", args=[d.pk])).status_code, 200)
        self.client.force_login(self.autre)
        self.assertEqual(self.client.get(reverse("document", args=[d.pk])).status_code, 403)
        autre_jalon = Jalon.objects.create(project=Projet.objects.create(
            owner=self.coordo, title="X", domain="TIC, numérique, IA et données"), title="J", due=dt.date.today())
        self.client.force_login(self.porteur)
        r = self.client.post(reverse("commentaire_ajout", args=[self.p.pk]), {"texte": "x", "jalon_id": autre_jalon.pk})
        self.assertEqual(r.status_code, 404)

    def test_commentaires_et_calendrier(self):
        self.client.force_login(self.porteur)
        self.client.post(reverse("commentaire_ajout", args=[self.p.pk]), {"texte": "Preuve envoyée", "jalon_id": self.j.pk})
        self.client.force_login(self.mentor)
        self.client.post(reverse("commentaire_ajout", args=[self.p.pk]), {"texte": "Bien reçu"})
        self.assertEqual(Commentaire.objects.count(), 2)
        r = self.client.get(reverse("projet", args=[self.p.pk]))
        self.assertContains(r, "Preuve envoyée")
        self.assertContains(r, "Bien reçu")
        r = self.client.get(reverse("calendrier") + "?m=" + self.j.due.strftime("%Y-%m"))
        self.assertContains(r, "Prototype")
        self.client.force_login(self.autre)
        self.assertNotContains(self.client.get(reverse("calendrier") + "?m=" + self.j.due.strftime("%Y-%m")), "Prototype")
        self.assertEqual(self.client.get(reverse("calendrier") + "?m=n'importe").status_code, 200)

    def test_compte_rendu_genere_puis_valide(self):
        self.client.force_login(self.mentor)
        with mock.patch.dict(os.environ):
            os.environ.pop("ANTHROPIC_API_KEY", None)
            r = self.client.post(reverse("compte_rendu", args=[self.p.pk]), {"action": "generer"})
        self.assertContains(r, "mode démonstration")
        self.assertContains(r, "Prototype : en retard")
        self.assertEqual(CompteRendu.objects.count(), 0)
        self.client.post(reverse("compte_rendu", args=[self.p.pk]),
                         {"action": "valider", "texte": "Compte rendu relu", "ia_used": "1"})
        cr = CompteRendu.objects.get()
        self.assertTrue(cr.ia_used)
        self.client.force_login(self.porteur)
        self.assertEqual(self.client.get(reverse("compte_rendu_pdf", args=[cr.pk])).content[:4], b"%PDF")
        self.assertEqual(self.client.get(reverse("compte_rendu", args=[self.p.pk])).status_code, 403)

    def test_tableau_de_bord_jalons_exclusifs(self):
        Jalon.objects.create(project=self.p, title="A", due=dt.date.today() + dt.timedelta(days=5))
        Jalon.objects.create(project=self.p, title="B", due=dt.date.today() + dt.timedelta(days=5), status="en_cours")
        Jalon.objects.create(project=self.p, title="C", due=dt.date.today(), status="valide")
        self.client.force_login(self.coordo)
        r = self.client.get(reverse("dashboard"))
        # validés, en cours, à faire, en retard : chaque jalon dans une seule catégorie
        self.assertEqual(r.context["chart_data"]["jalons"]["data"], [1, 1, 1, 1])
        self.assertEqual(r.context["s"]["st"], [("Accepté", 1, 100)])
        wb_bytes = self.client.get(reverse("export_xlsx")).content
        from openpyxl import load_workbook
        self.assertEqual(load_workbook(io.BytesIO(wb_bytes)).sheetnames, ["Projets", "Jalons", "Suivis 3-6 mois"])


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class AuditTests(TestCase):
    """Correctifs de l'audit : connexion limitée, mot de passe, cohérence, cohortes, quotas IA, mise en production."""

    def setUp(self):
        cache.clear()

        def mk(e, r):
            return User.objects.create_user(username=e, email=e, password="Passw0rd!", first_name=e, role=r)
        self.porteur, self.comite = mk("p@x.sn", "porteur"), mk("c@x.sn", "comite")
        self.coordo, self.mentor = mk("co@x.sn", "coordo"), mk("m@x.sn", "mentor")

    def test_connexion_bloquee_apres_5_echecs(self):
        for _ in range(5):
            self.client.post(reverse("login"), {"username": "p@x.sn", "password": "mauvais"})
        r = self.client.post(reverse("login"), {"username": "p@x.sn", "password": "Passw0rd!"})
        self.assertEqual(r.status_code, 429)
        self.assertContains(r, "Trop de tentatives", status_code=429)
        cache.clear()
        r = self.client.post(reverse("login"), {"username": "P@x.sn ", "password": "Passw0rd!"})
        self.assertEqual(r.status_code, 302)

    def test_changement_mot_de_passe(self):
        self.client.force_login(self.porteur)
        self.client.post(reverse("mot_de_passe"), {"old_password": "Passw0rd!", "new_password1": "Baobab-Diamniadio-26",
                                                   "new_password2": "Baobab-Diamniadio-26"})
        self.porteur.refresh_from_db()
        self.assertTrue(self.porteur.check_password("Baobab-Diamniadio-26"))

    def test_refus_retire_le_mentor(self):
        p = Projet.objects.create(owner=self.porteur, title="T", domain="TIC, numérique, IA et données",
                                  status="accepte", mentor=self.mentor)
        self.client.force_login(self.comite)
        self.client.post(reverse("evaluer", args=[p.pk]), {"innovation": 1, "faisabilite": 1, "impact": 1, "equipe": 1,
                                                           "decision": "refuse", "motivation": "Non viable"})
        p.refresh_from_db()
        self.assertEqual((p.status, p.mentor), ("refuse", None))

    def test_cohorte_par_defaut_et_modification(self):
        p = Projet.objects.create(owner=self.porteur, title="T", domain="TIC, numérique, IA et données")
        self.assertEqual(p.cohort, cohorte_courante())
        self.client.force_login(self.porteur)
        self.assertEqual(self.client.post(reverse("cohorte", args=[p.pk]), {"cohort": "2027-A"}).status_code, 403)
        self.client.force_login(self.coordo)
        self.client.post(reverse("cohorte", args=[p.pk]), {"cohort": "n'importe quoi"})
        self.client.post(reverse("cohorte", args=[p.pk]), {"cohort": "2027-A"})
        p.refresh_from_db()
        self.assertEqual(p.cohort, "2027-A")

    @override_settings(IA_APPELS_PAR_HEURE=2)
    def test_quota_ia(self):
        self.client.force_login(self.porteur)
        with mock.patch.dict(os.environ):
            os.environ.pop("ANTHROPIC_API_KEY", None)
            for _ in range(2):
                self.assertContains(self.client.post(reverse("ia"), {"texte": "Idée", "mode": "pitch"}), "Problème : Idée")
            r = self.client.post(reverse("ia"), {"texte": "Idée", "mode": "pitch"})
        self.assertContains(r, "Limite d&#x27;utilisation")
        self.assertNotContains(r, "Problème : Idée")

    def test_init_coordo_et_seed_demo(self):
        from django.core.management import call_command
        env = {"CIAT_ADMIN_EMAIL": "Direction@ciat.sn", "CIAT_ADMIN_PASSWORD": "Baobab-Diamniadio-26"}
        with mock.patch.dict(os.environ, env):
            call_command("init_coordo", stdout=io.StringIO())
            call_command("init_coordo", stdout=io.StringIO())  # idempotent
        admin = User.objects.get(username="direction@ciat.sn")
        self.assertTrue(admin.is_superuser and admin.role == "coordo")
        # porteur inscrit avec l'e-mail du coordonnateur : promu, et son mot de passe est remplacé
        with mock.patch.dict(os.environ, {"CIAT_ADMIN_EMAIL": "p@x.sn", "CIAT_ADMIN_PASSWORD": "Baobab-Diamniadio-26"}):
            call_command("init_coordo", stdout=io.StringIO())
            self.porteur.refresh_from_db()
            self.assertTrue(self.porteur.is_superuser and self.porteur.role == "coordo")
            self.assertTrue(self.porteur.check_password("Baobab-Diamniadio-26"))
            self.porteur.set_password("Nouveau-Mot-De-Passe-26")
            self.porteur.save()
            call_command("init_coordo", stdout=io.StringIO())  # déjà coordonnateur : mot de passe conservé
            self.porteur.refresh_from_db()
            self.assertTrue(self.porteur.check_password("Nouveau-Mot-De-Passe-26"))
        with mock.patch.dict(os.environ, {"CIAT_ADMIN_EMAIL": "p@x.sn", "CIAT_ADMIN_PASSWORD": " Baobab-Diamniadio-26\n",
                                          "CIAT_ADMIN_RESET": "1"}):
            call_command("init_coordo", stdout=io.StringIO())  # récupération : mot de passe réimposé, espaces ignorés
            self.porteur.refresh_from_db()
            self.assertTrue(self.porteur.check_password("Baobab-Diamniadio-26"))
        with mock.patch.dict(os.environ):
            os.environ.pop("DEMO_PASSWORD", None)
            call_command("seed", stdout=io.StringIO())  # DEBUG=False en test : refus sans mot de passe
        self.assertFalse(User.objects.filter(username="coordo@ciat.sn").exists())
        with mock.patch.dict(os.environ, {"DEMO_PASSWORD": "Jury-CIAT-2026!"}):
            call_command("seed", stdout=io.StringIO())
        demo = User.objects.get(username="coordo@ciat.sn")
        self.assertTrue(demo.check_password("Jury-CIAT-2026!"))
        self.assertFalse(demo.is_staff or demo.is_superuser)
        self.assertEqual(Projet.objects.filter(owner__username__endswith="@exemple.sn").exists(), True)
        call_command("seed", "--supprimer", stdout=io.StringIO())
        self.assertFalse(User.objects.filter(username__in=["coordo@ciat.sn", "porteur3@exemple.sn"]).exists())
        self.assertTrue(User.objects.filter(username="direction@ciat.sn").exists())
        self.assertTrue(User.objects.filter(username="p@x.sn").exists())

    def test_pages_erreur(self):
        self.client.force_login(self.porteur)
        self.assertContains(self.client.get(reverse("dashboard")), "Accès refusé", status_code=403)
        self.assertContains(self.client.get("/nexiste-pas/"), "Page introuvable", status_code=404)

    def test_reglages_production(self):
        import subprocess
        import sys
        code = ("import django, os; os.environ['DJANGO_SETTINGS_MODULE']='config.settings'; django.setup(); "
                "from django.conf import settings as s; "
                "print(s.DATABASES['default']['ENGINE'], s.STORAGES['default']['BACKEND'], "
                "s.STORAGES['default']['OPTIONS']['default_acl'])")
        env = dict(os.environ, DJANGO_DEBUG="0", DJANGO_SECRET_KEY="x" * 50,
                   DATABASE_URL="postgres://u:p@ep-test.eu-west-2.aws.neon.tech/ciat", R2_BUCKET="ciat",
                   R2_ENDPOINT="https://compte.r2.cloudflarestorage.com", R2_ACCESS_KEY_ID="a", R2_SECRET_ACCESS_KEY="b")
        out = subprocess.run([sys.executable, "-c", code], env=env, capture_output=True, text=True)
        self.assertEqual(out.stdout.split(), ["django.db.backends.postgresql", "storages.backends.s3.S3Storage", "None"])
        env.pop("DJANGO_SECRET_KEY")
        out = subprocess.run([sys.executable, "-c", code], env=env, capture_output=True, text=True)
        self.assertIn("DJANGO_SECRET_KEY est obligatoire", out.stderr)
