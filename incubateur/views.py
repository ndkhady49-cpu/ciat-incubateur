"""Vues : authentification, projets, évaluation, suivi (jalons, preuves, commentaires, suivis 3/6 mois),
calendrier, tableau de bord, comptes rendus, exports, assistant IA."""
import calendar
import datetime as dt
import io
import os
from functools import wraps

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import LoginView, PasswordChangeView
from django.core.cache import cache
from django.core.exceptions import PermissionDenied
from django.db.models import Count, Prefetch, Q
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse_lazy
from django.views.decorators.http import require_POST

from .forms import (AccompagnementForm, CohorteForm, CompteRenduForm, DocumentForm, EmailAuthForm, EvaluationForm,
                    JalonForm, ProfilForm, ProjetForm, RegisterForm, SuiviForm)
from .ia import CONSIGNES, compte_rendu_ia, suggestion_ia
from .models import DOMAINES, STATUTS, Commentaire, CompteRendu, Document, Jalon, Projet, Suivi, User


# ---------- droits ----------
def role_required(*roles):
    """Exige une connexion et, si précisé, l'un des rôles donnés (sinon 403)."""
    def deco(view):
        @wraps(view)
        def w(request, *a, **k):
            if roles and request.user.role not in roles:
                raise PermissionDenied
            return view(request, *a, **k)
        return login_required(w)
    return deco


def projet_autorise(request, pk):
    """Retourne le projet si l'utilisateur a le droit de le voir, sinon 403."""
    p = get_object_or_404(Projet.objects.select_related("owner", "mentor"), pk=pk)
    u = request.user
    if (u.role in ("coordo", "comite") or (u.role == "porteur" and p.owner_id == u.id)
            or (u.role == "mentor" and p.mentor_id == u.id)):
        return p
    raise PermissionDenied


# ---------- limites anti-abus ----------
def ip_client(request):
    """Adresse IP du visiteur (Render place l'IP réelle en tête de X-Forwarded-For)."""
    fwd = request.META.get("HTTP_X_FORWARDED_FOR", "")
    return fwd.split(",")[0].strip() if fwd else request.META.get("REMOTE_ADDR", "")


def compter(cle, duree):
    """Incrémente un compteur qui expire après « duree » secondes ; retourne la nouvelle valeur."""
    cache.add(cle, 0, duree)
    try:
        return cache.incr(cle)
    except ValueError:  # expiré entre add et incr
        cache.set(cle, 1, duree)
        return 1


def quota_ia(user):
    """True si l'utilisateur peut encore appeler l'IA cette heure-ci (sinon False)."""
    return compter("ia:%d" % user.pk, 3600) <= settings.IA_APPELS_PAR_HEURE


class ConnexionView(LoginView):
    """Connexion par e-mail, bloquée 15 minutes après 5 échecs pour un même e-mail ou une même adresse IP."""
    template_name = "auth.html"
    authentication_form = EmailAuthForm
    extra_context = {"mode": "login"}

    def cles(self):
        email = self.request.POST.get("username", "").strip().lower()
        return ["login:ip:" + ip_client(self.request), "login:email:" + email]

    def post(self, request, *args, **kwargs):
        if any(cache.get(k, 0) >= settings.LOGIN_ECHECS_MAX for k in self.cles()):
            return self.render_to_response(self.get_context_data(form=self.authentication_form(request),
                                                                 bloque=True), status=429)
        return super().post(request, *args, **kwargs)

    def form_invalid(self, form):
        for k in self.cles():
            compter(k, settings.LOGIN_BLOCAGE_SECONDES)
        return super().form_invalid(form)

    def form_valid(self, form):
        cache.delete_many(self.cles())
        return super().form_valid(form)


class MotDePasseView(PasswordChangeView):
    template_name = "auth.html"
    extra_context = {"mode": "password"}
    success_url = reverse_lazy("profil")

    def form_valid(self, form):
        messages.success(self.request, "Mot de passe modifié.")
        return super().form_valid(form)


# ---------- accès ----------
def accueil(request):
    if request.user.is_authenticated:
        return redirect("dashboard" if request.user.role in ("coordo", "comite") else "projets")
    nb_projets = Projet.objects.count()
    nb_mentors = User.objects.filter(role="mentor").count()
    decides = Projet.objects.filter(status__in=("accepte", "refuse")).count()
    acceptes = Projet.objects.filter(status="accepte").count()
    taux_sel = int(100 * acceptes / decides) if decides else 0
    return render(request, "accueil.html", {
        "nb_projets": nb_projets, "nb_mentors": nb_mentors, "taux_sel": taux_sel})


def register(request):
    form = RegisterForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        login(request, form.save())
        messages.success(request, "Compte créé. Bienvenue !")
        return redirect("accueil")
    return render(request, "auth.html", {"form": form, "mode": "register"})


def projets_visibles(u):
    """Projets que l'utilisateur peut consulter."""
    qs = Projet.objects.select_related("owner")
    if u.role == "porteur":
        return qs.filter(owner=u)
    if u.role == "mentor":
        return qs.filter(mentor=u)
    return qs


def jalon_du_projet(request, p):
    """Jalon désigné par jalon_id dans le formulaire (doit appartenir au projet), ou None."""
    jid = request.POST.get("jalon_id", "")
    if not jid:
        return None
    if not jid.isdigit():
        raise PermissionDenied
    return get_object_or_404(Jalon, pk=int(jid), project=p)


# ---------- projets ----------
@role_required()
def projets(request):
    return render(request, "projets.html", {"rows": projets_visibles(request.user)})


@role_required("porteur")
def projet_form(request, pk=None):
    p = projet_autorise(request, pk) if pk else None
    if p and not p.editable:
        raise PermissionDenied
    form = ProjetForm(request.POST or None, request.FILES or None, instance=p)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        if not p:
            obj.owner = request.user
        else:
            obj.status = "soumis"
        obj.save()
        pc, manque = obj.completude()
        messages.success(request, "Dossier enregistré (%d %% complet)." % pc
                         + (" Manque : " + ", ".join(manque) if manque else ""))
        return redirect("projet", pk=obj.pk)
    return render(request, "projet_form.html", {"form": form, "p": p})


@role_required()
def projet(request, pk):
    p = projet_autorise(request, pk)
    pc, manque = p.completude()
    jalons = p.jalons.prefetch_related(
        Prefetch("preuves", queryset=Document.objects.select_related("auteur")),
        Prefetch("commentaires", queryset=Commentaire.objects.select_related("auteur")))
    faits = set(p.suivis.values_list("periode", flat=True))
    return render(request, "projet.html", {
        "p": p, "pc": pc, "manque": manque, "jalons": jalons,
        "evals": p.evaluations.select_related("evaluator"),
        "mentors": User.objects.filter(role="mentor"),
        "documents": p.documents.filter(jalon=None).select_related("auteur"),
        "commentaires": p.commentaires.filter(jalon=None).select_related("auteur"),
        "suivis": p.suivis.select_related("auteur"),
        "echeances": [(m, d, m in faits) for m, d in p.echeances_suivi()],
        "comptes_rendus": p.comptes_rendus.select_related("auteur"),
        "eval_form": EvaluationForm(), "jalon_statuts": Jalon.STATUTS,
        "doc_form": DocumentForm(), "acc_form": AccompagnementForm(instance=p), "suivi_form": SuiviForm(),
        "cohorte_form": CohorteForm(instance=p),
        "peut_suivre": request.user.role in ("coordo", "mentor") and p.status == "accepte",
        "peut_deposer": request.user.role in ("porteur", "mentor", "coordo"), "today": dt.date.today()})


@role_required("comite", "coordo")
@require_POST
def evaluer(request, pk):
    p = projet_autorise(request, pk)
    form = EvaluationForm(request.POST)
    if form.is_valid():
        ev = form.save(commit=False)
        ev.project, ev.evaluator = p, request.user
        ev.save()
        p.status = ev.decision
        if ev.decision != "accepte" and p.mentor_id:
            p.mentor = None  # un projet non accepté n'est plus accompagné
            messages.info(request, "Le mentor a été retiré : le projet n'est plus accepté.")
        p.save(update_fields=["status", "mentor"])
        messages.success(request, "Décision enregistrée.")
    else:
        messages.error(request, "Notes de 0 à 5, décision et motivation obligatoires.")
    return redirect("projet", pk=pk)


@role_required("coordo")
@require_POST
def affecter(request, pk):
    p = projet_autorise(request, pk)
    mentor = User.objects.filter(pk=request.POST.get("mentor_id") or 0, role="mentor").first()
    if p.status != "accepte" or not mentor:
        messages.error(request, "Un mentor ne peut être affecté qu'à un projet accepté.")
    else:
        p.mentor = mentor
        p.save(update_fields=["mentor"])
        messages.success(request, "Mentor affecté.")
    return redirect("projet", pk=pk)


@role_required("coordo", "mentor")
@require_POST
def jalon_ajout(request, pk):
    p = projet_autorise(request, pk)
    form = JalonForm(request.POST)
    if form.is_valid():
        j = form.save(commit=False)
        j.project = p
        j.save()
    else:
        messages.error(request, "Titre et date valides obligatoires.")
    return redirect("projet", pk=pk)


@role_required("coordo", "mentor")
@require_POST
def jalon_maj(request, pk):
    j = get_object_or_404(Jalon, pk=pk)
    projet_autorise(request, j.project_id)
    if request.POST.get("status") in dict(Jalon.STATUTS):
        j.status = request.POST["status"]
        j.save(update_fields=["status"])
    return redirect("projet", pk=j.project_id)


@role_required("coordo")
@require_POST
def cohorte(request, pk):
    p = projet_autorise(request, pk)
    form = CohorteForm(request.POST, instance=p)
    if form.is_valid():
        form.save()
        messages.success(request, "Cohorte mise à jour.")
    else:
        messages.error(request, "Cohorte invalide (format attendu : 2026-A).")
    return redirect("projet", pk=pk)


@role_required("coordo", "mentor")
@require_POST
def accompagnement(request, pk):
    """Objectifs d'accompagnement et date de fin d'incubation (déclenche les suivis à 3 et 6 mois)."""
    p = projet_autorise(request, pk)
    form = AccompagnementForm(request.POST, instance=p)
    if p.status != "accepte":
        messages.error(request, "Les objectifs se définissent pour un projet accepté.")
    elif form.is_valid():
        form.save()
        messages.success(request, "Objectifs d'accompagnement enregistrés.")
    else:
        messages.error(request, "Date de fin d'incubation invalide.")
    return redirect("projet", pk=pk)


@role_required("porteur", "mentor", "coordo")
@require_POST
def document_ajout(request, pk):
    """Ajout d'un document au dossier ou d'une preuve sur un jalon (si jalon_id est fourni)."""
    p = projet_autorise(request, pk)
    jalon = jalon_du_projet(request, p)
    form = DocumentForm(request.POST, request.FILES)
    if form.is_valid():
        d = form.save(commit=False)
        d.project, d.jalon, d.auteur = p, jalon, request.user
        d.save()
        messages.success(request, "Preuve ajoutée au jalon." if jalon else "Document ajouté au dossier.")
    else:
        messages.error(request, "Document refusé : " + " ".join(e for errs in form.errors.values() for e in errs))
    return redirect("projet", pk=pk)


@role_required()
def document(request, pk):
    d = get_object_or_404(Document, pk=pk)
    projet_autorise(request, d.project_id)
    return FileResponse(d.fichier.open("rb"), as_attachment=True, filename=os.path.basename(d.fichier.name))


@role_required()
@require_POST
def commentaire_ajout(request, pk):
    p = projet_autorise(request, pk)
    jalon = jalon_du_projet(request, p)
    texte = request.POST.get("texte", "").strip()
    if texte:
        Commentaire.objects.create(project=p, jalon=jalon, auteur=request.user, texte=texte[:1000])
    else:
        messages.error(request, "Le commentaire est vide.")
    return redirect("projet", pk=pk)


@role_required("coordo", "mentor")
@require_POST
def suivi_ajout(request, pk):
    """Saisie du suivi post-incubation à 3 ou 6 mois (un seul par période)."""
    p = projet_autorise(request, pk)
    form = SuiviForm(request.POST)
    if p.status != "accepte":
        messages.error(request, "Le suivi post-incubation concerne les projets acceptés.")
    elif not form.is_valid():
        messages.error(request, "Période, situation et nombre d'emplois valides obligatoires.")
    elif p.suivis.filter(periode=form.cleaned_data["periode"]).exists():
        messages.error(request, "Le suivi à %s mois est déjà saisi." % form.cleaned_data["periode"])
    else:
        s = form.save(commit=False)
        s.project, s.auteur = p, request.user
        s.save()
        messages.success(request, "Suivi à %d mois enregistré." % s.periode)
    return redirect("projet", pk=pk)


@role_required()
def fichier(request, pk):
    p = projet_autorise(request, pk)
    if not p.file:
        raise Http404
    return FileResponse(p.file.open("rb"), as_attachment=True, filename=os.path.basename(p.file.name))


def pdf_response(titre, sections, nom_fichier):
    """PDF simple (titre puis sections « libellé : texte »), avec saut de page automatique."""
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setTitle(titre)
    y = 800

    def ligne(texte, police, x, pas):
        nonlocal y
        if y < 60:
            c.showPage()
            y = 800
        c.setFont(police, 11)
        c.drawString(x, y, texte)
        y -= pas

    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, titre)
    y -= 30
    for lib, val in sections:
        ligne(lib + " :", "Helvetica-Bold", 50, 15)
        for paragraphe in (val or "-").splitlines():
            for morceau in simpleSplit(paragraphe, "Helvetica", 11, 490) or [""]:
                ligne(morceau, "Helvetica", 60, 14)
        y -= 8
    c.save()
    r = HttpResponse(buf.getvalue(), content_type="application/pdf")
    r["Content-Disposition"] = 'inline; filename="%s"' % nom_fichier
    return r


@role_required()
def fiche_pdf(request, pk):
    p = projet_autorise(request, pk)
    return pdf_response("CIAT Incubateur - Fiche projet", [
        ("Titre", p.title), ("Porteur", p.owner.nom()), ("Domaine", p.domain), ("Cohorte", p.cohort),
        ("Statut", p.get_status_display()), ("Problème", p.problem), ("Solution", p.solution),
        ("Mentor", p.mentor.nom() if p.mentor else ""), ("Objectifs d'accompagnement", p.objectifs)],
        "fiche_projet_%d.pdf" % p.pk)


# ---------- comptes rendus (rédaction assistée par IA, validation humaine) ----------
@role_required("coordo", "mentor")
def compte_rendu(request, pk):
    """GET : formulaire vide. POST action=generer : brouillon IA. POST action=valider : enregistrement."""
    p = projet_autorise(request, pk)
    src = None
    action = request.POST.get("action")
    if request.method == "POST" and action == "valider":
        form = CompteRenduForm(request.POST)
        if form.is_valid():
            cr = form.save(commit=False)
            cr.project, cr.auteur = p, request.user
            cr.ia_used = request.POST.get("ia_used") == "1"
            cr.save()
            messages.success(request, "Compte rendu enregistré.")
            return redirect("projet", pk=pk)
    elif request.method == "POST" and action == "generer":
        if quota_ia(request.user):
            texte, src = compte_rendu_ia(p)
            form = CompteRenduForm(initial={"texte": texte})
        else:
            messages.error(request, "Limite d'utilisation de l'IA atteinte : réessayez dans une heure.")
            form = CompteRenduForm()
    else:
        form = CompteRenduForm()
    return render(request, "compte_rendu.html", {"p": p, "form": form, "src": src,
                                                 "ia_used": src is not None or request.POST.get("ia_used") == "1"})


@role_required()
def compte_rendu_pdf(request, pk):
    cr = get_object_or_404(CompteRendu.objects.select_related("auteur"), pk=pk)
    p = projet_autorise(request, cr.project_id)
    sections = [("Projet", "%s (%s, cohorte %s)" % (p.title, p.domain, p.cohort)), ("Porteur", p.owner.nom()),
                ("Mentor", p.mentor.nom() if p.mentor else "non affecté"),
                ("Rédigé par", "%s le %s" % (cr.auteur.nom(), cr.created.strftime("%d/%m/%Y"))),
                ("Compte rendu", cr.texte)]
    if cr.ia_used:
        sections.append(("Mention", "Rédigé avec l'aide d'une IA, relu et validé par l'auteur."))
    return pdf_response("CIAT Incubateur - Compte rendu de suivi", sections, "compte_rendu_%d.pdf" % cr.pk)


# ---------- calendrier ----------
MOIS = ["", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août", "Septembre", "Octobre",
        "Novembre", "Décembre"]


@role_required()
def calendrier(request):
    """Calendrier mensuel des échéances de jalons et des suivis à 3/6 mois des projets visibles."""
    today = dt.date.today()
    try:
        an, mois = [int(x) for x in request.GET.get("m", "").split("-")]
        debut = dt.date(an, mois, 1)
    except ValueError:
        debut = today.replace(day=1)
    fin = debut.replace(day=calendar.monthrange(debut.year, debut.month)[1])
    projets = projets_visibles(request.user)
    evts = {}
    for j in Jalon.objects.filter(project__in=projets, due__range=(debut, fin)).select_related("project"):
        evts.setdefault(j.due, []).append({"titre": j.title, "projet": j.project,
                                           "classe": "ok" if j.status == "valide" else "ko" if j.en_retard else ""})
    for p in projets.filter(fin_incubation__isnull=False).prefetch_related("suivis"):
        faits = {s.periode for s in p.suivis.all()}
        for m, d in p.echeances_suivi():
            if debut <= d <= fin:
                evts.setdefault(d, []).append({"titre": "Suivi à %d mois" % m, "projet": p,
                                               "classe": "ok" if m in faits else "ko" if d < today else "suivi"})
    semaines = [[{"date": d, "dans_mois": d.month == debut.month, "evts": evts.get(d, [])} for d in sem]
                for sem in calendar.Calendar().monthdatescalendar(debut.year, debut.month)]
    prec = (debut - dt.timedelta(days=1)).replace(day=1)
    suiv = fin + dt.timedelta(days=1)
    return render(request, "calendrier.html", {
        "semaines": semaines, "titre": "%s %d" % (MOIS[debut.month], debut.year), "today": today,
        "prec": prec.strftime("%Y-%m"), "suiv": suiv.strftime("%Y-%m"), "liste": sorted(evts.items()),
        "jours": ["Lun", "Mar", "Mer", "Jeu", "Ven", "Sam", "Dim"]})


# ---------- tableau de bord ----------
def _barres(qs, col):
    rows = list(qs.order_by().values(col).annotate(n=Count("id"), ok=Count("id", filter=Q(status="accepte"))).order_by("-n"))
    mx = max([r["n"] for r in rows], default=1)
    return [{"k": r[col] or "non renseigné", "n": r["n"], "ok": r["ok"], "pct": int(100 * r["n"] / mx)} for r in rows]


@role_required("coordo", "comite")
def dashboard(request):
    f = {k: request.GET.get(k, "") for k in ("cohort", "domain", "status")}
    qs = Projet.objects.filter(**{k: v for k, v in f.items() if v})
    st = {r["status"]: r["n"] for r in qs.order_by().values("status").annotate(n=Count("id"))}
    decides = st.get("accepte", 0) + st.get("refuse", 0)
    today = dt.date.today()
    jalons = Jalon.objects.filter(project__in=qs)
    nj = jalons.count()
    # Catégories exclusives : un jalon non validé dont l'échéance est passée compte uniquement « en retard »
    jalons_valides = jalons.filter(status="valide").count()
    retard_qs = jalons.exclude(status="valide").filter(due__lt=today)
    jalons_retard = retard_qs.count()
    a_jour = jalons.exclude(status="valide").filter(due__gte=today)
    jalons_encours = a_jour.filter(status="en_cours").count()
    jalons_afaire = a_jour.filter(status="a_faire").count()
    suivis_dus = [(p, m, d) for p in qs.filter(status="accepte", fin_incubation__lte=today).select_related("owner")
                  for m, d in p.suivis_dus()]
    suivis = Suivi.objects.filter(project__in=qs)
    total = qs.count()
    stats = {"total": total,
             "acceptes": st.get("accepte", 0),
             "refuses": st.get("refuse", 0),
             "soumis": st.get("soumis", 0),
             "a_completer": st.get("a_completer", 0),
             "taux": int(100 * st.get("accepte", 0) / decides) if decides else 0,
             "retards": jalons_retard,
             "avancement": int(100 * jalons_valides / nj) if nj else 0,
             "jalons_total": nj, "jalons_valides": jalons_valides, "jalons_encours": jalons_encours,
             "jalons_afaire": jalons_afaire, "suivis_dus": len(suivis_dus), "suivis_faits": suivis.count(),
             "emplois": sum(suivis.values_list("emplois", flat=True)),
             "ia_count": qs.filter(ia_used=True).count(),
             "mentors_actifs": User.objects.filter(role="mentor", projets_suivis__in=qs).distinct().count(),
             "st": [(dict(STATUTS)[k], v, round(100 * v / total)) for k, v in st.items()]}
    # Données JSON pour Chart.js
    par_domaine = _barres(qs, "domain")
    par_cohorte = _barres(qs, "cohort")
    genre = _barres(qs, "owner__genre")
    chart_data = {
        "statuts": {"labels": [lib for lib, n, pct in stats["st"]], "data": [n for lib, n, pct in stats["st"]]},
        "jalons": {"labels": ["Validés", "En cours", "À faire", "En retard"],
                   "data": [jalons_valides, jalons_encours, jalons_afaire, jalons_retard]},
        "domaines": {"labels": [r["k"][:25] for r in par_domaine], "data": [r["n"] for r in par_domaine],
                     "ok": [r["ok"] for r in par_domaine]},
        "cohortes": {"labels": [r["k"] for r in par_cohorte], "data": [r["n"] for r in par_cohorte],
                     "ok": [r["ok"] for r in par_cohorte]},
        "genre": {"labels": [r["k"] for r in genre], "data": [r["n"] for r in genre]},
    }
    return render(request, "dashboard.html", {
        "f": f, "s": stats, "domaines": DOMAINES, "statuts": STATUTS,
        "cohortes": Projet.objects.order_by("cohort").values_list("cohort", flat=True).distinct(),
        "par_domaine": par_domaine, "par_cohorte": par_cohorte, "genre": genre,
        "retards": retard_qs.select_related("project").order_by("due")[:10], "suivis_dus": suivis_dus[:10],
        "chart_data": chart_data})


@role_required("coordo")
def export_xlsx(request):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Projets"
    ws.append(["ID", "Titre", "Porteur", "Domaine", "Cohorte", "Statut", "Mentor", "Complétude %"])
    for p in Projet.objects.select_related("owner", "mentor").order_by("id"):
        ws.append([p.id, p.title, p.owner.nom(), p.domain, p.cohort, p.get_status_display(),
                   p.mentor.nom() if p.mentor else "", p.completude()[0]])
    ws = wb.create_sheet("Jalons")
    ws.append(["Projet", "Jalon", "Échéance", "Statut", "En retard", "Preuves"])
    for j in Jalon.objects.select_related("project").annotate(np=Count("preuves")).order_by("project_id", "due"):
        ws.append([j.project.title, j.title, j.due, j.get_status_display(), "oui" if j.en_retard else "non", j.np])
    ws = wb.create_sheet("Suivis 3-6 mois")
    ws.append(["Projet", "Période (mois)", "Situation", "Emplois créés", "Chiffre d'affaires (FCFA)", "Observations"])
    for s in Suivi.objects.select_related("project").order_by("project_id", "periode"):
        ws.append([s.project.title, s.periode, s.get_situation_display(), s.emplois, s.chiffre_affaires, s.observations])
    buf = io.BytesIO()
    wb.save(buf)
    r = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r["Content-Disposition"] = 'attachment; filename="projets_ciat.xlsx"'
    return r


# ---------- assistant IA (suggestion + validation humaine) ----------
@role_required("porteur")
def ia(request):
    sug = src = None
    mode, texte = request.POST.get("mode", "reformuler"), request.POST.get("texte", "").strip()
    if request.method == "POST" and texte and mode in CONSIGNES:
        if quota_ia(request.user):
            sug, src = suggestion_ia(mode, texte[:3000])
        else:
            messages.error(request, "Limite d'utilisation de l'IA atteinte : réessayez dans une heure.")
    mes = Projet.objects.filter(owner=request.user, status__in=("soumis", "a_completer"))
    return render(request, "ia.html", {"sug": sug, "src": src, "texte": texte, "mode": mode,
                                       "mes": mes, "modes": list(CONSIGNES)})


@role_required("porteur")
@require_POST
def ia_accepter(request):
    pid = request.POST.get("project_id", "")
    if not pid.isdigit():
        raise PermissionDenied
    p = projet_autorise(request, int(pid))
    champ = request.POST.get("champ")
    if champ not in ("problem", "solution") or not p.editable:
        raise PermissionDenied
    setattr(p, champ, request.POST.get("texte", "").strip())
    p.ia_used = True
    p.save()
    messages.success(request, "Texte validé et enregistré (rédigé avec l'aide d'une IA).")
    return redirect("projet", pk=p.pk)


# ---------- profil ----------
@role_required()
def profil(request):
    form = ProfilForm(request.POST or None, user=request.user)
    if request.method == "POST" and form.is_valid():
        request.user.first_name = form.cleaned_data["name"]
        request.user.genre = form.cleaned_data["genre"]
        request.user.save(update_fields=["first_name", "genre"])
        messages.success(request, "Profil mis à jour.")
        return redirect("profil")
    return render(request, "profil.html", {"form": form})
